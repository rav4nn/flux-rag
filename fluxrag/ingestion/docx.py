"""DOCX parser using python-docx."""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument

from fluxrag.core.schema import Document
from fluxrag.ingestion.base import AbstractParser


class DocxParser(AbstractParser):
    """Parses DOCX files, extracting paragraphs and tables."""

    def parse(self, source_path: str, **kwargs: object) -> list[Document]:
        path = Path(source_path)
        docx = DocxDocument(str(path))
        parts: list[str] = []

        for para in docx.paragraphs:
            if para.text.strip():
                parts.append(para.text.strip())

        for table in docx.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))

        full_text = "\n\n".join(parts)
        if not full_text.strip():
            return []

        return [
            Document(
                text=full_text,
                metadata={
                    "source_type": "docx",
                    "source_path": str(path),
                    "title": path.stem,
                },
            )
        ]

    def supported_extensions(self) -> list[str]:
        return [".docx"]
