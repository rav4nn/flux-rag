"""Tests for DOCX parser."""

from pathlib import Path

import pytest
from docx import Document as DocxDocument

from fluxrag.ingestion.docx import DocxParser


parser = DocxParser()


@pytest.fixture
def sample_docx(tmp_dir: Path) -> Path:
    p = tmp_dir / "sample.docx"
    doc = DocxDocument()
    doc.add_paragraph("Water temperature affects extraction rate.")
    doc.add_paragraph("Use 93°C for light roasts and 88°C for dark roasts.")
    doc.save(str(p))
    return p


@pytest.fixture
def empty_docx(tmp_dir: Path) -> Path:
    p = tmp_dir / "empty.docx"
    doc = DocxDocument()
    doc.save(str(p))
    return p


@pytest.fixture
def docx_with_table(tmp_dir: Path) -> Path:
    p = tmp_dir / "table.docx"
    doc = DocxDocument()
    doc.add_paragraph("Brew method comparison:")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Method"
    table.cell(0, 1).text = "Grind"
    table.cell(1, 0).text = "Espresso"
    table.cell(1, 1).text = "Fine"
    doc.save(str(p))
    return p


def test_parse_docx(sample_docx: Path) -> None:
    docs = parser.parse(str(sample_docx))
    assert len(docs) == 1
    assert "temperature" in docs[0].text
    assert docs[0].metadata["source_type"] == "docx"


def test_empty_docx_returns_nothing(empty_docx: Path) -> None:
    docs = parser.parse(str(empty_docx))
    assert docs == []


def test_extracts_tables(docx_with_table: Path) -> None:
    docs = parser.parse(str(docx_with_table))
    assert len(docs) == 1
    assert "Espresso" in docs[0].text
    assert "Fine" in docs[0].text
